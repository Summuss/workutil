"""Main seam tests for Excel export.

Calls the export endpoint to retrieve bytes, then loads them back with openpyxl
to assert sheet names, cell values, formatting, image count and anchor rows
(design.md §6 F5, Testing Decisions).
"""

import io
from pathlib import Path
from typing import Any
from urllib.parse import unquote

import openpyxl
from fastapi.testclient import TestClient
from PIL import Image as PILImage

from app.modules.evidence.layout import LayoutSettings

from .conftest import data_url


def create_evidence(client: TestClient, title: str = "受注一覧の絞り込み修正") -> int:
    res = client.post("/api/evidence", json={"title": title})
    assert res.status_code == 201
    return int(res.json()["id"])


def add_case(client: TestClient, evidence_id: int, name: str) -> int:
    res = client.post(f"/api/evidence/{evidence_id}/cases", json={"name": name})
    assert res.status_code == 201
    return int(res.json()["id"])


def add_text_block(
    client: TestClient,
    evidence_id: int,
    case_id: int,
    text: str,
    label: str | None = None,
) -> int:
    payload: dict[str, str | None] = {"kind": "text", "text": text}
    if label is not None:
        payload["label"] = label
    at = f"/api/evidence/{evidence_id}/cases/{case_id}/blocks"
    res = client.post(at, json=payload)
    assert res.status_code == 201
    return int(res.json()["id"])


def add_table_block(
    client: TestClient,
    evidence_id: int,
    case_id: int,
    tsv: str,
    label: str | None = None,
) -> int:
    payload: dict[str, str | None] = {"kind": "paste", "text": tsv}
    if label is not None:
        payload["label"] = label
    at = f"/api/evidence/{evidence_id}/cases/{case_id}/blocks"
    res = client.post(at, json=payload)
    assert res.status_code == 201
    assert res.json()["kind"] == "table"
    return int(res.json()["id"])


def add_image_block(
    client: TestClient,
    evidence_id: int,
    case_id: int,
    image_bytes: bytes,
    label: str | None = None,
) -> int:
    payload = {
        "kind": "image",
        "image": {"data": data_url(image_bytes), "filename": "screenshot.png"},
        "label": label,
    }
    res = client.post(
        f"/api/evidence/{evidence_id}/cases/{case_id}/blocks",
        json=payload,
    )
    assert res.status_code == 201, res.text
    assert res.json()["kind"] == "image"
    return int(res.json()["id"])


def make_png(width: int, height: int, color: str = "blue") -> bytes:
    im = PILImage.new("RGB", (width, height), color=color)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def test_export_nonexistent_evidence_returns_404(client: TestClient) -> None:
    res = client.get("/api/evidence/999999/export")
    assert res.status_code == 404


def test_export_evidence_with_no_cases_returns_422(client: TestClient) -> None:
    evidence_id = create_evidence(client, "空のEvidence")
    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 422


def test_export_case_with_no_blocks_produces_empty_sheet(client: TestClient) -> None:
    evidence_id = create_evidence(client, "ブロックなし")
    add_case(client, evidence_id, "用例1")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200
    assert (
        res.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["用例1"]
    ws = wb["用例1"]
    # Empty case has no cell values
    assert ws.cell(row=1, column=1).value is None


def test_export_multiple_sheets_match_case_names_and_order(
    client: TestClient,
) -> None:
    evidence_id = create_evidence(client, "複数用例")
    add_case(client, evidence_id, "1")
    add_case(client, evidence_id, "2~5")
    add_case(client, evidence_id, "特別対応")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["1", "2~5", "特別対応"]


def test_export_filename_and_rfc5987_content_disposition(
    client: TestClient,
) -> None:
    evidence_id = create_evidence(client, r'受注/一覧:改*定?"<>|')
    add_case(client, evidence_id, "1")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    disposition = res.headers["content-disposition"]
    assert "filename*=UTF-8''" in disposition
    encoded_part = disposition.split("filename*=UTF-8''")[1]
    decoded_filename = unquote(encoded_part)
    assert decoded_filename == "エビデンス_受注_一覧_改_定_____.xlsx"


def test_export_text_blocks_and_labels_layout(client: TestClient) -> None:
    evidence_id = create_evidence(client, "文字ブロック排版")
    case_id = add_case(client, evidence_id, "1")

    add_text_block(client, evidence_id, case_id, "事前ログ\n行2", label="事前準備")
    add_text_block(client, evidence_id, case_id, "事後確認テキスト")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    # Row 1: label "事前準備", bold, @
    c1 = ws.cell(row=1, column=1)
    assert c1.value == "事前準備"
    assert c1.font.bold is True
    assert c1.number_format == "@"

    # Row 2: text "事前ログ\n行2", @
    c2 = ws.cell(row=2, column=1)
    assert c2.value == "事前ログ\n行2"
    assert c2.number_format == "@"

    # Row 3: empty row separating blocks
    c3 = ws.cell(row=3, column=1)
    assert c3.value is None

    # Row 4: text "事後確認テキスト" (no label), @
    c4 = ws.cell(row=4, column=1)
    assert c4.value == "事後確認テキスト"
    assert c4.font.bold is not True
    assert c4.number_format == "@"


def test_guardrail_cell_with_007_remains_string(client: TestClient) -> None:
    """Guardrail: values like '007', dates, and IDs must stay strings under '@'."""
    evidence_id = create_evidence(client, "007保持テスト")
    case_id = add_case(client, evidence_id, "1")

    tsv = "code\tdate\tid\n007\t2024-01-01\t123456789012345678"
    add_table_block(client, evidence_id, case_id, tsv)

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    # Data row is row 2
    cell_007 = ws.cell(row=2, column=1)
    assert cell_007.value == "007"
    assert isinstance(cell_007.value, str)
    assert cell_007.number_format == "@"

    cell_date = ws.cell(row=2, column=2)
    assert cell_date.value == "2024-01-01"
    assert isinstance(cell_date.value, str)
    assert cell_date.number_format == "@"

    cell_id = ws.cell(row=2, column=3)
    assert cell_id.value == "123456789012345678"
    assert isinstance(cell_id.value, str)
    assert cell_id.number_format == "@"


def test_table_block_header_and_border_styling(client: TestClient) -> None:
    evidence_id = create_evidence(client, "テーブル装飾")
    case_id = add_case(client, evidence_id, "1")

    # Table 1: with header (default)
    tsv1 = "colA\tcolB\nval1\tval2"
    add_table_block(client, evidence_id, case_id, tsv1)

    # Table 2: header flipped to False
    tsv2 = "num1\tnum2\n100\t200"
    block2_id = add_table_block(client, evidence_id, case_id, tsv2)
    patch_res = client.put(
        f"/api/evidence/{evidence_id}/cases/{case_id}/blocks/{block2_id}/header",
        json={"has_header": False},
    )
    assert patch_res.status_code == 200

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    # Table 1 starts at row 1
    # Row 1 is header: bold + #87e7ad + thin border
    h1 = ws.cell(row=1, column=1)
    assert h1.value == "colA"
    assert h1.font.bold is True
    assert h1.fill.fill_type == "solid"
    assert h1.fill.start_color.rgb == "0087E7AD" or h1.fill.start_color.rgb == "87E7AD"
    assert h1.border.left.style == "thin"

    # Row 2 is data: not bold, no fill, thin border
    d1 = ws.cell(row=2, column=1)
    assert d1.value == "val1"
    assert d1.font.bold is not True
    assert d1.fill.fill_type is None
    assert d1.border.left.style == "thin"

    # Row 3 is empty row separating Table 1 and Table 2
    assert ws.cell(row=3, column=1).value is None

    # Table 2 starts at row 4
    # has_header=False: row 4 is not bold and has no fill
    h2 = ws.cell(row=4, column=1)
    assert h2.value == "num1"
    assert h2.font.bold is not True
    assert h2.fill.fill_type is None
    assert h2.border.left.style == "thin"

    d2 = ws.cell(row=5, column=1)
    assert d2.value == "100"
    assert d2.font.bold is not True
    assert d2.fill.fill_type is None
    assert d2.border.left.style == "thin"


def test_guardrail_1080p_image_block_reserves_rows_and_next_block_is_below(
    client: TestClient,
) -> None:
    """Guardrail: 1080px image must reserve ceil(1080 / 20) = 54 rows.

    The subsequent block MUST start strictly below the reserved rows.
    An implementation that set row height to image height would fail here
    because Excel row height max is 409 pt (= 545 px).
    """
    evidence_id = create_evidence(client, "1080p画像排版")
    case_id = add_case(client, evidence_id, "1")

    # 800 x 1080 image (height 1080 px, width <= 900 so height is unscaled)
    img_bytes = make_png(800, 1080)
    add_image_block(client, evidence_id, case_id, img_bytes, label="画面1")
    add_text_block(client, evidence_id, case_id, "画像の下のテキスト")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    # Row 1 is label "画面1"
    assert ws.cell(row=1, column=1).value == "画面1"

    # Image is anchored at row 2
    ws_any: Any = ws
    images = ws_any._images
    assert len(images) == 1
    img = images[0]
    # 0-indexed row in openpyxl drawing anchor: row 2 is index 1
    image_anchor_row = img.anchor._from.row + 1
    assert image_anchor_row == 2

    # 1080 px reserves 54 rows (rows 2 through 55).
    # With 1 empty row (row 56), the next text block is at row 57.
    # It must be strictly > image_anchor_row + 53.
    text_cell = None
    text_row = None
    for r in range(50, 70):
        val = ws.cell(row=r, column=1).value
        if val == "画像の下のテキスト":
            text_cell = ws.cell(row=r, column=1)
            text_row = r
            break

    assert text_cell is not None
    assert text_row is not None
    # Must be at or below row 56 (anchor row 2 + 54 rows reserved)
    assert text_row >= image_anchor_row + 54


def test_sheet_column_dimensions_are_set(client: TestClient) -> None:
    evidence_id = create_evidence(client, "列幅テスト")
    case_id = add_case(client, evidence_id, "1")
    tsv = "col1\tcol2\tcol3\na\tb\tc"
    add_table_block(client, evidence_id, case_id, tsv)

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    settings = LayoutSettings()
    assert ws.column_dimensions["A"].width == settings.column_width
    assert ws.column_dimensions["B"].width == settings.column_width
    assert ws.column_dimensions["C"].width == settings.column_width


def test_guardrail_text_starting_with_equals_is_not_read_as_a_formula(
    client: TestClient,
) -> None:
    """Guardrail: '@' alone does not stop openpyxl reading a leading '=' as a
    formula. A value like '=1+1' must round-trip as the literal string, not
    silently evaluate to '2' — the same promise '@' makes for '007'."""
    evidence_id = create_evidence(client, "数式化テスト")
    case_id = add_case(client, evidence_id, "1")
    add_text_block(client, evidence_id, case_id, "=1+1")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    cell = wb["1"].cell(row=1, column=1)
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_guardrail_table_cell_starting_with_equals_is_not_read_as_a_formula(
    client: TestClient,
) -> None:
    evidence_id = create_evidence(client, "テーブル数式化テスト")
    case_id = add_case(client, evidence_id, "1")
    add_table_block(client, evidence_id, case_id, "expr\tnote\n=SUM(1,2)\tok")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    cell = wb["1"].cell(row=2, column=1)
    assert cell.value == "=SUM(1,2)"
    assert cell.data_type == "s"


def test_control_characters_in_a_paste_do_not_break_the_export(
    client: TestClient,
) -> None:
    """Guardrail: OOXML forbids control characters other than tab/CR/LF, and
    openpyxl raises on them. A pasted terminal log carrying an ANSI escape
    must still export, with the illegal byte stripped rather than a 500."""
    evidence_id = create_evidence(client, "制御文字テスト")
    case_id = add_case(client, evidence_id, "1")
    add_table_block(client, evidence_id, case_id, "col\tval\n\x1b[31mred\x1b[0m\tok")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    cell = wb["1"].cell(row=2, column=1)
    assert cell.value == "[31mred[0m"


def test_export_refuses_when_an_image_file_is_missing_from_disk(
    client: TestClient, data_dir: Path
) -> None:
    """Guardrail: a screenshot referenced by a block but absent on disk must
    refuse the whole export (422), not silently swap in placeholder text —
    a partial workbook is not a deliverable (design.md §6 F5)."""
    evidence_id = create_evidence(client, "画像欠落テスト")
    case_id = add_case(client, evidence_id, "1")
    block_id = add_image_block(client, evidence_id, case_id, make_png(10, 10))

    case = client.get(f"/api/evidence/{evidence_id}/cases/{case_id}").json()
    image_url = next(b["image_url"] for b in case["blocks"] if b["id"] == block_id)
    image_name = image_url.rsplit("/", 1)[-1]
    (data_dir / "images" / "evidence" / str(evidence_id) / image_name).unlink()

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 422


def test_export_font_name_and_size_applied_everywhere(client: TestClient) -> None:
    evidence_id = create_evidence(client, "フォントテスト")
    case_id = add_case(client, evidence_id, "1")
    add_text_block(client, evidence_id, case_id, "本文テキスト", label="見出し")
    add_table_block(client, evidence_id, case_id, "col1\tcol2\nval1\tval2")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    # Label (row 1)
    label_cell = ws.cell(row=1, column=1)
    assert label_cell.value == "見出し"
    assert label_cell.font.name == "游ゴシック"
    assert label_cell.font.size == 11
    assert label_cell.font.bold is True

    # Text block (row 2)
    text_cell = ws.cell(row=2, column=1)
    assert text_cell.value == "本文テキスト"
    assert text_cell.font.name == "游ゴシック"
    assert text_cell.font.size == 11

    # Row 3 is empty separator

    # Table header (row 4)
    th_cell = ws.cell(row=4, column=1)
    assert th_cell.value == "col1"
    assert th_cell.font.name == "游ゴシック"
    assert th_cell.font.size == 11
    assert th_cell.font.bold is True

    # Table data (row 5)
    td_cell = ws.cell(row=5, column=1)
    assert td_cell.value == "val1"
    assert td_cell.font.name == "游ゴシック"
    assert td_cell.font.size == 11
    assert td_cell.font.bold is not True


def test_table_header_alignment_wrap_and_center_data_row_default(
    client: TestClient,
) -> None:
    evidence_id = create_evidence(client, "配置テスト")
    case_id = add_case(client, evidence_id, "1")
    add_table_block(client, evidence_id, case_id, "long_col_name\tother_col\nval\t1")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    # Header cell: wrap_text=True and vertical="center"
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.alignment.wrap_text is True
    assert header_cell.alignment.vertical == "center"

    # Data cell: default alignment (wrap_text not set, vertical not set)
    data_cell = ws.cell(row=2, column=1)
    assert data_cell.alignment.wrap_text is not True
    assert data_cell.alignment.vertical is None


def test_null_literal_exact_match_gray_color_and_variants_stay_black(
    client: TestClient,
) -> None:
    evidence_id = create_evidence(client, "NULL色テスト")
    case_id = add_case(client, evidence_id, "1")

    # 5 test values in order:
    # 1. Exact ≪ NULL ≫ (U+226A + ' NULL ' + U+226B)
    # 2. ≪NULL≫ (no spaces)
    # 3. « NULL » (U+00AB / U+00BB Guillemets)
    # 4. NULL (plain)
    # 5. 　≪ NULL ≫　 (leading/trailing whitespace)
    v1 = "\u226a NULL \u226b"
    v2 = "\u226aNULL\u226b"
    v3 = "\u00ab NULL \u00bb"
    v4 = "NULL"
    v5 = "　\u226a NULL \u226b　"

    tsv = f"head1\thead2\thead3\thead4\thead5\n{v1}\t{v2}\t{v3}\t{v4}\t{v5}"
    add_table_block(client, evidence_id, case_id, tsv)

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    # Header cells must not be gray
    for col in range(1, 6):
        h_cell = ws.cell(row=1, column=col)
        assert h_cell.font.color is None or h_cell.font.color.rgb not in (
            "808080",
            "00808080",
        )

    # Col 1: exact match -> gray #808080
    c1 = ws.cell(row=2, column=1)
    assert c1.value == v1
    assert c1.font.color is not None
    assert c1.font.color.rgb in ("808080", "00808080")

    # Col 2-5: variants must NOT turn gray
    for col in (2, 3, 4, 5):
        c = ws.cell(row=2, column=col)
        assert c.font.color is None or c.font.color.rgb not in ("808080", "00808080")


def test_text_block_null_literal_does_not_turn_gray(client: TestClient) -> None:
    evidence_id = create_evidence(client, "TextブロックNULL色テスト")
    case_id = add_case(client, evidence_id, "1")
    add_text_block(client, evidence_id, case_id, "\u226a NULL \u226b")

    res = client.get(f"/api/evidence/{evidence_id}/export")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["1"]

    cell = ws.cell(row=1, column=1)
    assert cell.value == "\u226a NULL \u226b"
    assert cell.font.color is None or cell.font.color.rgb not in ("808080", "00808080")
